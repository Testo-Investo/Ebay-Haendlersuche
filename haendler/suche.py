"""Je Produktgruppe die größten gewerblichen eBay-Händler finden.

„Größter Händler“ = die meisten passenden Angebote in der Gruppe, bei Gleichstand die
meisten Bewertungen. Bestehende Kontakte (Mailinglisten des Ghostwriters) und Neobution
selbst fallen raus, damit nur Neukunden in der Liste stehen.
"""

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .ebay import Ebay, impressum
from .gruppen import Gruppe

TOP = 20
FREEMAIL = {"gmail.com", "googlemail.com", "hotmail.com", "outlook.com", "yahoo.com", "ymail.com", "icloud.com",
            "gmx.de", "gmx.net", "web.de", "t-online.de", "live.com", "aol.com", "freenet.de"}
EIGENE = ("neobution",)


@dataclass
class Haendler:
    verkaeufer: str
    bewertungen: int = 0
    positiv: str = ""
    angebote: set = field(default_factory=set)
    beispiel_id: str = ""
    beispiel_url: str = ""
    kontakt: dict = field(default_factory=dict)


def bestehende_kontakte(ordner: Path | None) -> tuple[set[str], set[str]]:
    """(E-Mail-Adressen, Firmen-Domains) aus den Mailinglisten – diese Händler kennt Marc schon."""
    adressen, domains = set(), set()
    if not ordner or not ordner.exists():
        return adressen, domains
    for datei in ordner.glob("*.xlsx"):
        for zeile in load_workbook(datei, read_only=True).worksheets[0].iter_rows(values_only=True):
            for zelle in zeile:
                text = str(zelle or "").strip().strip("'\"<>;,").lower()
                if "@" in text and "." in text.split("@")[-1]:
                    adressen.add(text)
                    if text.split("@")[1] not in FREEMAIL:
                        domains.add(text.split("@")[1])
    return adressen, domains


def _schon_kunde(kontakt: dict, verkaeufer: str, adressen: set, domains: set) -> bool:
    email = kontakt.get("email", "").lower()
    if any(e in (verkaeufer + kontakt.get("firma", "") + email).lower() for e in EIGENE):
        return True
    return bool(email) and (email in adressen or email.split("@")[-1] in domains)


def suche_gruppe(ebay: Ebay, gruppe: Gruppe, cache: dict, adressen: set, domains: set) -> list[Haendler]:
    gefunden: dict[str, Haendler] = {}
    anfragen = [{"gtin": ean} for ean in gruppe.suchbegriffe()] + [{"q": gruppe.name}]
    for anfrage in anfragen:
        for item in ebay.suche(**anfrage):
            s = item.get("seller") or {}
            name = s.get("username")
            if not name:
                continue
            h = gefunden.setdefault(name, Haendler(name))
            h.bewertungen = int(s.get("feedbackScore") or 0)
            h.positiv = str(s.get("feedbackPercentage") or "")
            h.angebote.add(item.get("itemId"))
            if not h.beispiel_id:
                h.beispiel_id, h.beispiel_url = item.get("itemId", ""), item.get("itemWebUrl", "")
    rangliste = sorted(gefunden.values(), key=lambda h: (len(h.angebote), h.bewertungen), reverse=True)
    out = []
    for h in rangliste:
        if len(out) >= TOP:
            break
        if h.verkaeufer not in cache:  # Impressum je Händler nur einmal abrufen
            cache[h.verkaeufer] = impressum(ebay.artikel(h.beispiel_id)) if h.beispiel_id else {}
        h.kontakt = cache[h.verkaeufer]
        if _schon_kunde(h.kontakt, h.verkaeufer, adressen, domains):
            continue
        out.append(h)
    return out
