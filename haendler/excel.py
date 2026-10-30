"""Ergebnis als Excel: Blatt „Händler“ mit Marcs Spalten, Blatt „Details“ mit eBay-Angaben."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SPALTEN = ["Produktgruppe", "Name", "Vorname", "Firma", "Adresse", "E-Mail-Adresse", "Telefonnummer"]
DETAILS = ["Produktgruppe", "Rang", "Firma", "eBay-Verkäufer", "Bewertungen", "Positiv %",
           "Passende Angebote", "Beispiel-Angebot"]
KOPF = PatternFill("solid", fgColor="1F4E78")


def _blatt(ws, kopf: list[str], zeilen: list[list], breiten: list[int]) -> None:
    ws.append(kopf)
    for zelle in ws[1]:
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = KOPF
        zelle.alignment = Alignment(vertical="center")
    for z in zeilen:
        ws.append(z)
    for i, b in enumerate(breiten, 1):
        ws.column_dimensions[get_column_letter(i)].width = b
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def schreibe(ergebnis: list[tuple[str, list]], pfad: Path) -> Path:
    """ergebnis: [(Produktgruppe, [Haendler, …]), …] in der Reihenfolge der Gruppen."""
    wb = Workbook()
    haendler, details = [], []
    for gruppe, liste in ergebnis:
        for rang, h in enumerate(liste, 1):
            k = h.kontakt
            firma = k.get("firma") or h.verkaeufer
            haendler.append([gruppe, k.get("name", ""), k.get("vorname", ""), firma, k.get("adresse", ""),
                             k.get("email", ""), k.get("telefon", "")])
            details.append([gruppe, rang, firma, h.verkaeufer, h.bewertungen, h.positiv, len(h.angebote),
                            h.beispiel_url])
    _blatt(wb.active, SPALTEN, haendler, [30, 18, 16, 34, 48, 34, 20])
    wb.active.title = "Händler"
    _blatt(wb.create_sheet("Details"), DETAILS, details, [30, 6, 34, 24, 12, 10, 18, 60])
    pfad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pfad)
    return pfad
